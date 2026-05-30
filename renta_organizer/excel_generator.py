"""Generates the fiscal summary Excel with tabs by category."""
from __future__ import annotations

import logging
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from renta_organizer.classifier import ClassifiedExpense
from renta_organizer.cross_reference import CrossReferenceReport
from renta_organizer.pdf_extractor import InvoiceData

logger = logging.getLogger(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
AMOUNT_FORMAT = '#,##0.00 €'

CATEGORY_LABELS = {
    "sanitario": "Sanitario",
    "formacion": "Formacion",
    "donaciones": "Donaciones",
    "seguros": "Seguros",
    "suministros": "Suministros",
    "transporte": "Transporte",
    "vivienda": "Vivienda",
    "actividad_profesional": "Actividad profesional",
    "potencialmente_deducible": "Potencial deducible",
    "no_deducible": "No deducible",
    "ingreso": "Ingresos",
}


def generate_excel(
    all_expenses: list[ClassifiedExpense],
    xref: CrossReferenceReport,
    invoices: list[InvoiceData],
    year: int,
    output_dir: Path,
) -> Path:
    """Generate the fiscal summary Excel."""
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"resumen_fiscal_{year}.xlsx"

    wb = openpyxl.Workbook()

    _create_summary_sheet(wb, all_expenses, year)
    _create_category_sheets(wb, all_expenses)
    _create_all_movements_sheet(wb, all_expenses)
    _create_no_invoice_sheet(wb, xref)
    _create_unmatched_invoices_sheet(wb, xref)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(path)
    logger.info("Excel generado: %s", path)

    _organize_pdf_folders(all_expenses, invoices, output_dir)

    return path


def _style_header(ws, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)


def _create_summary_sheet(wb, expenses: list[ClassifiedExpense], year: int) -> None:
    ws = wb.active
    ws.title = "Resumen"

    ws.append([f"Resumen Fiscal {year}"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([])

    headers = ["Categoria", "Num. Gastos", "Total (€)", "Con factura", "Sin factura"]
    ws.append(headers)
    _style_header(ws, len(headers))
    ws.insert_rows(3, 0)

    categories_order = [
        "sanitario", "formacion", "donaciones", "seguros",
        "suministros", "transporte", "vivienda", "actividad_profesional",
        "potencialmente_deducible", "no_deducible",
    ]

    total_all = 0
    for cat in categories_order:
        cat_expenses = [e for e in expenses if e.category == cat]
        if not cat_expenses:
            continue
        total = sum(abs(e.transaction.amount) for e in cat_expenses)
        with_invoice = sum(1 for e in cat_expenses if e.invoice_path)
        without_invoice = len(cat_expenses) - with_invoice
        total_all += total

        row = [
            CATEGORY_LABELS.get(cat, cat),
            len(cat_expenses),
            total,
            with_invoice,
            without_invoice,
        ]
        ws.append(row)
        ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT

    ws.append([])
    ws.append(["TOTAL", sum(1 for e in expenses if e.category != "ingreso"), total_all])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)

    income = [e for e in expenses if e.category == "ingreso"]
    if income:
        ws.append([])
        total_income = sum(e.transaction.amount for e in income)
        ws.append(["Ingresos totales", len(income), total_income])
        ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT

    _auto_width(ws)


def _create_category_sheets(wb, expenses: list[ClassifiedExpense]) -> None:
    categories_with_data = {}
    for e in expenses:
        if e.category in ("no_deducible", "ingreso"):
            continue
        categories_with_data.setdefault(e.category, []).append(e)

    for cat, cat_expenses in categories_with_data.items():
        label = CATEGORY_LABELS.get(cat, cat)[:25]
        ws = wb.create_sheet(title=label)

        headers = ["Fecha", "Concepto", "Importe (€)", "Banco", "Factura", "Emisor factura", "Motivo clasificacion", "Confianza"]
        ws.append(headers)
        _style_header(ws, len(headers))

        for e in sorted(cat_expenses, key=lambda x: x.transaction.date):
            ws.append([
                e.transaction.date_str,
                e.transaction.description[:80],
                abs(e.transaction.amount),
                e.transaction.bank,
                Path(e.invoice_path).name if e.invoice_path else "",
                e.invoice_issuer,
                e.reason,
                f"{e.confidence:.0%}",
            ])
            ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT

        total_row = ws.max_row + 1
        ws.cell(row=total_row, column=2, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=sum(abs(e.transaction.amount) for e in cat_expenses))
        ws.cell(row=total_row, column=3).number_format = AMOUNT_FORMAT
        ws.cell(row=total_row, column=3).font = Font(bold=True)

        _auto_width(ws)


def _create_all_movements_sheet(wb, expenses: list[ClassifiedExpense]) -> None:
    ws = wb.create_sheet(title="Todos movimientos")
    headers = ["Fecha", "Concepto", "Importe (€)", "Banco", "Categoria", "Motivo", "Factura"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for e in sorted(expenses, key=lambda x: x.transaction.date):
        ws.append([
            e.transaction.date_str,
            e.transaction.description[:80],
            e.transaction.amount,
            e.transaction.bank,
            CATEGORY_LABELS.get(e.category, e.category),
            e.reason,
            Path(e.invoice_path).name if e.invoice_path else "",
        ])
        ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT

    _auto_width(ws)


def _create_no_invoice_sheet(wb, xref: CrossReferenceReport) -> None:
    if not xref.expenses_without_invoice:
        return

    ws = wb.create_sheet(title="Sin factura")
    headers = ["Fecha", "Concepto", "Importe (€)", "Banco", "Categoria", "Motivo"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for mr in sorted(xref.expenses_without_invoice, key=lambda x: x.expense.transaction.date):
        e = mr.expense
        ws.append([
            e.transaction.date_str,
            e.transaction.description[:80],
            abs(e.transaction.amount),
            e.transaction.bank,
            CATEGORY_LABELS.get(e.category, e.category),
            e.reason,
        ])
        ws.cell(row=ws.max_row, column=3).number_format = AMOUNT_FORMAT

    _auto_width(ws)


def _create_unmatched_invoices_sheet(wb, xref: CrossReferenceReport) -> None:
    if not xref.invoices_without_match:
        return

    ws = wb.create_sheet(title="Facturas sin match")
    headers = ["Archivo", "Emisor", "CIF", "Fecha", "Importe (€)", "Descripcion"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for inv in xref.invoices_without_match:
        ws.append([
            Path(inv.file_path).name,
            inv.issuer[:60],
            inv.cif,
            inv.date,
            inv.total_amount,
            inv.description[:80],
        ])
        if inv.total_amount:
            ws.cell(row=ws.max_row, column=5).number_format = AMOUNT_FORMAT

    _auto_width(ws)


def _organize_pdf_folders(
    expenses: list[ClassifiedExpense],
    invoices: list[InvoiceData],
    output_dir: Path,
) -> None:
    """Copy matched PDFs into organized category folders."""
    facturas_dir = output_dir / "facturas"

    for e in expenses:
        if not e.invoice_path:
            continue
        src = Path(e.invoice_path)
        if not src.exists():
            continue

        cat_dir = facturas_dir / e.category
        cat_dir.mkdir(parents=True, exist_ok=True)
        dest = cat_dir / src.name
        if not dest.exists():
            shutil.copy2(src, dest)

    unmatched_dir = facturas_dir / "sin_clasificar"
    matched_paths = {e.invoice_path for e in expenses if e.invoice_path}
    for inv in invoices:
        if inv.file_path not in matched_paths:
            src = Path(inv.file_path)
            if src.exists():
                unmatched_dir.mkdir(parents=True, exist_ok=True)
                dest = unmatched_dir / src.name
                if not dest.exists():
                    shutil.copy2(src, dest)
